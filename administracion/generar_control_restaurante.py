# -*- coding: utf-8 -*-
"""
Genera el libro "Control_Restaurante_La_Cremeria.xlsx":
inventario, compras, recetas / fichas técnicas, costeo y precios, ventas,
mermas, planilla, gastos, cierre de caja, estado de resultados,
punto de equilibrio, flujo de caja a 12 meses y dashboard.

Uso:  python3 administracion/generar_control_restaurante.py
Todos los cálculos quedan como fórmulas de Excel (nada calculado en Python).
Los datos de ejemplo son supuestos didácticos: reemplácelos por los reales.
"""
import datetime as dt
import os

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "Control_Restaurante_La_Cremeria.xlsx")

# --------------------------------------------------------------------------
# Estilos (convención de la capacitación: celdas a llenar en AZUL,
# celdas calculadas en GRIS)
# --------------------------------------------------------------------------
FONT = "Arial"
C_DARK = "1F3864"
C_TEAL = "2E8B8B"
F_TITLE = Font(name=FONT, size=16, bold=True, color=C_TEAL)
F_SUB = Font(name=FONT, size=9, italic=True, color="595959")
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_BODY = Font(name=FONT, size=10, color="000000")
F_IN = Font(name=FONT, size=10, color="0000FF")          # entrada
F_LINK = Font(name=FONT, size=10, color="008000")        # viene de otra hoja
F_BOLD = Font(name=FONT, size=10, bold=True)
F_SEC = Font(name=FONT, size=12, bold=True, color=C_DARK)

FILL_HDR = PatternFill("solid", fgColor=C_DARK)
FILL_HDR_IN = PatternFill("solid", fgColor="2F75B5")
FILL_IN = PatternFill("solid", fgColor="DDEBF7")        # azul claro = llenar
FILL_CALC = PatternFill("solid", fgColor="F2F2F2")      # gris = automático
FILL_KEY = PatternFill("solid", fgColor="FFFF00")       # amarillo = clave
FILL_TOT = PatternFill("solid", fgColor="FCE4D6")
FILL_SEC = PatternFill("solid", fgColor="E2EFDA")

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="top", wrap_text=True)

CRC = '"₡"#,##0;[Red]-"₡"#,##0;"-"'
CRC2 = '"₡"#,##0.00;[Red]-"₡"#,##0.00;"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
NUM = '#,##0;[Red]-#,##0;"-"'
NUM1 = '#,##0.0;[Red]-#,##0.0;"-"'
DATE = "dd/mm/yyyy"
MONTH = "mmm-yyyy"

# Filas de datos de cada registro
HDR = 5
FIRST = 6
LAST = {
    "INSUMOS": 205, "COMPRAS": 1005, "RECETAS": 605, "PRODUCTOS": 105,
    "VENTAS": 2005, "MERMAS": 505, "PLANILLA": 505, "GASTOS": 1005,
    "CAJA_DIARIA": 405,
}


def R(sheet, col):
    """Rango absoluto de una columna completa de datos de un registro."""
    return f"{sheet}!${col}${FIRST}:${col}${LAST[sheet]}"


def LOOK(sheet, ret_col, key_cell, key_col="A"):
    return f"INDEX({R(sheet, ret_col)},MATCH({key_cell},{R(sheet, key_col)},0))"


wb = Workbook()


def name(nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def title(ws, text, sub, sub2=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    ws["A2"] = sub
    ws["A2"].font = F_SUB
    if sub2:
        ws["A3"] = sub2
        ws["A3"].font = F_SUB
    ws.sheet_view.showGridLines = False


def style(c, font=F_BODY, fill=None, fmt=None, align=None, border=True):
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = align
    if border:
        c.border = BORDER


def register(ws, sheet, cols, examples=(), filt=True):
    """Crea un registro (tabla) con encabezados en fila 5 y datos desde la 6.
    cols: lista de dicts {h: encabezado, w: ancho, kind: 'in'|'calc',
          f: fórmula con {r} (sólo calc), fmt, note}
    """
    last = LAST[sheet]
    for j, c in enumerate(cols, start=1):
        L = get_column_letter(j)
        h = ws.cell(row=HDR, column=j, value=c["h"])
        style(h, F_HDR, FILL_HDR_IN if c["kind"] == "in" else FILL_HDR,
              align=CENTER)
        if c.get("note"):
            h.comment = Comment(c["note"], "Control")
        ws.column_dimensions[L].width = c.get("w", 12)
        for r in range(FIRST, last + 1):
            cell = ws.cell(row=r, column=j)
            if c["kind"] == "in":
                style(cell, F_IN, FILL_IN, c.get("fmt"))
            else:
                cell.value = "=" + c["f"].format(r=r)
                style(cell, F_BODY, FILL_CALC, c.get("fmt"))
    ws.row_dimensions[HDR].height = 42
    for i, row in enumerate(examples):
        for j, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(row=FIRST + i, column=j, value=v)
    ws.freeze_panes = ws.cell(row=FIRST, column=3)
    if filt:
        ws.auto_filter.ref = f"A{HDR}:{get_column_letter(len(cols))}{last}"


def dv_list(ws, formula, rng):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = "Elija un valor de la lista"
    dv.showErrorMessage = False
    ws.add_data_validation(dv)
    dv.add(rng)


def status_colors(ws, rng, first_cell, bad, warn=(), good=("OK",)):
    red = PatternFill("solid", fgColor="F8CBAD")
    yel = PatternFill("solid", fgColor="FFE699")
    grn = PatternFill("solid", fgColor="C6EFCE")
    for v in bad:
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'{first_cell}="{v}"'], fill=red,
                             font=Font(name=FONT, bold=True, color="9C0006")))
    for v in warn:
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'{first_cell}="{v}"'], fill=yel))
    for v in good:
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'{first_cell}="{v}"'], fill=grn))


def blank_guard(cell, expr, empty='""'):
    return f'IF({cell}="",{empty},{expr})'


# ==========================================================================
# 1. INICIO
# ==========================================================================
ws = wb.active
ws.title = "INICIO"
title(ws, "CONTROL ADMINISTRATIVO – RESTAURANTE / HELADERÍA",
      "Inventario, costos, recetas, planilla, gastos, caja y reportes en un solo libro.")
ws.column_dimensions["A"].width = 4
ws.column_dimensions["B"].width = 26
ws.column_dimensions["C"].width = 95
ws["B4"] = "Negocio:"
ws["B4"].font = F_BOLD
ws["C4"] = "=NEGOCIO"
ws["C4"].font = Font(name=FONT, size=14, bold=True, color=C_DARK)

ws["B6"] = "CÓMO LEER LOS COLORES"
ws["B6"].font = F_SEC
leg = [
    (FILL_IN, F_IN, "Azul claro, letra azul", "Celda para LLENAR por usted (datos de entrada)."),
    (FILL_CALC, F_BODY, "Gris", "Se calcula sola. No la escriba (tiene fórmula)."),
    (FILL_CALC, F_LINK, "Letra verde", "Viene de otra hoja (enlace)."),
    (FILL_KEY, F_BOLD, "Amarillo", "Resultado o supuesto clave."),
]
for i, (fill, font, a, b) in enumerate(leg):
    r = 7 + i
    style(ws.cell(row=r, column=2, value=a), font, fill)
    style(ws.cell(row=r, column=3, value=b), F_BODY)

ws["B12"] = "ORDEN DE USO"
ws["B12"].font = F_SEC
steps = [
    ("CONFIG", "1 sola vez: nombre, año, % IVA, cargas sociales CCSS/INS, márgenes objetivo y listas desplegables."),
    ("INSUMOS", "Catálogo de materias primas y empaques: presentación de compra, precio, % aprovechamiento, stock mínimo e inicial. Aquí ve el INVENTARIO ACTUAL, su valor y qué reordenar."),
    ("COMPRAS", "Registre cada factura de proveedor. Suma al inventario y actualiza el último costo del insumo. Controle pendientes de pago."),
    ("RECETAS", "Una línea por ingrediente de cada producto (cantidad neta por porción o por lote). Calcula el costo de cada receta."),
    ("FICHA_TECNICA", "Elija un producto y vea su ficha técnica: ingredientes, costo, margen, precio sugerido y preparación."),
    ("PRODUCTOS", "Menú con precio de venta (con IVA). Calcula costo variable, contribución, food cost y precio sugerido."),
    ("VENTAS", "Registre ventas (por ticket o total diario por producto). Descuenta el inventario según las recetas."),
    ("MERMAS", "Mermas, vencidos, cortesías y ajustes de conteo físico. Afectan inventario y resultados."),
    ("PLANILLA", "Planilla mensual: salario, horas extra, CCSS obrera/patronal, INS, aguinaldo, vacaciones, cesantía."),
    ("GASTOS", "Gastos operativos (alquiler, luz, agua, internet, publicidad, contador, etc.)."),
    ("CAJA_DIARIA", "Cierre de caja: fondo, ventas por medio de pago, efectivo esperado vs contado, faltantes/sobrantes."),
    ("ESTADO_RESULTADOS", "Automático: ventas, costo de ventas, utilidad, planilla, gastos, indicadores e IVA por mes."),
    ("PUNTO_EQUILIBRIO", "Presupuesto mensual por producto, punto de equilibrio y calculadoras de precio (recargo vs margen)."),
    ("FLUJO_CAJA", "Proyección de 12 meses: entradas, pagos previstos, aguinaldo, reserva mínima y alertas de faltante."),
    ("DASHBOARD", "Resumen del mes elegido: KPIs, productos más rentables, insumos por reordenar y gráfico anual."),
]
for i, (s, d) in enumerate(steps):
    r = 13 + i
    c = ws.cell(row=r, column=2, value=s)
    c.hyperlink = f"#'{s}'!A1"
    style(c, Font(name=FONT, size=10, bold=True, color="0563C1", underline="single"))
    style(ws.cell(row=r, column=3, value=d), F_BODY, align=WRAP)
    ws.row_dimensions[r].height = 28

r = 13 + len(steps) + 1
ws.cell(row=r, column=2, value="REGLAS DE ORO").font = F_SEC
rules = [
    "Precios de venta con IVA en PRODUCTOS; los análisis usan el precio NETO (sin IVA), como en la capacitación.",
    "Las recetas usan cantidades NETAS aprovechables (g, ml o unidades). El % de aprovechamiento convierte a cantidad bruta comprada.",
    "La planilla fija va una sola vez (en PLANILLA), nunca dentro del costo de las recetas.",
    "Use el mismo período en todo y no duplique costos (p. ej., no meta en GASTOS lo que ya está en COMPRAS de insumos).",
    "Haga un conteo físico al menos semanal en INSUMOS (columna 'Conteo físico') para detectar faltantes.",
    "Los datos que trae el libro son EJEMPLOS ilustrativos: bórrelos o reemplácelos por los suyos. Verifique tasas CCSS/INS/renta vigentes con su contador.",
]
for i, t in enumerate(rules):
    c = ws.cell(row=r + 1 + i, column=2, value="•")
    c.alignment = Alignment(horizontal="right")
    style(ws.cell(row=r + 1 + i, column=3, value=t), F_BODY, align=WRAP, border=False)

# ==========================================================================
# 2. CONFIG
# ==========================================================================
cf = wb.create_sheet("CONFIG")
title(cf, "CONFIGURACIÓN Y PARÁMETROS",
      "Llene las celdas azules. Las demás hojas usan estos valores por su nombre (p. ej. IVA_VENTAS).")
cf.column_dimensions["A"].width = 3
cf.column_dimensions["B"].width = 44
cf.column_dimensions["C"].width = 16
cf.column_dimensions["D"].width = 50
for c, t in (("B", "Parámetro"), ("C", "Valor"), ("D", "Nota / fuente")):
    style(cf[f"{c}5"], F_HDR, FILL_HDR, align=CENTER)
    cf[f"{c}5"] = t
params = [
    ("Nombre del negocio", "La Cremería", None, "NEGOCIO", "Aparece en reportes."),
    ("Año del estado de resultados", 2026, "0", "ANIO", "Columnas ene–dic de ESTADO_RESULTADOS."),
    ("Mes a analizar en el DASHBOARD", dt.date(2026, 10, 1), MONTH, "MES_REPORTE", "Escriba el día 1 del mes."),
    ("IVA en ventas", 0.13, PCT, "IVA_VENTAS", "Tarifa general CR 13%. Se usa como valor por defecto en PRODUCTOS."),
    ("IVA promedio en compras de insumos", 0.05, PCT, "IVA_COMPRAS_PROM", "Supuesto para el flujo de caja (mezcla canasta básica 1% y 13%)."),
    ("CCSS cuota patronal", 0.265, PCT, "CCSS_PATRONAL", "Valor usado en la capacitación (26.50%). Verifique la tasa vigente."),
    ("CCSS cuota obrera (rebajo al trabajador)", 0.1067, PCT, "CCSS_OBRERO", "Supuesto. Verifique la tasa vigente."),
    ("Póliza riesgos del trabajo (INS)", 0.0185, PCT, "INS_RT", "Valor de la capacitación (1.85%). Depende de su póliza."),
    ("Provisión aguinaldo", 0.0833, PCT, "PROV_AGUINALDO", "1/12 del salario."),
    ("Provisión vacaciones", 0.0416, PCT, "PROV_VACACIONES", "Valor de la capacitación (4.16%)."),
    ("Provisión cesantía / despidos", 0.0533, PCT, "PROV_CESANTIA", "Valor de la capacitación (5.33%)."),
    ("Recargo hora extra (x salario/hora)", 1.5, "0.00", "RECARGO_HE", "Tiempo y medio."),
    ("Horas base por mes (salario/hora)", 240, "0", "HORAS_MES", "30 días x 8 horas."),
    ("Imprevistos en flujo de caja", 0.05, PCT, "IMPREVISTOS", "Valor de la capacitación (5%)."),
    ("Impuesto sobre la renta estimado", 0.15, PCT, "RENTA", "Supuesto de la capacitación. Persona jurídica PYME / física tiene tramos: confirme con su contador."),
    ("Margen de contribución objetivo", 0.65, PCT, "MARGEN_OBJ", "Para precio sugerido = costo ÷ (1 − margen)."),
    ("Food cost máximo aceptable", 0.35, PCT, "FOODCOST_MAX", "Si un producto lo supera se marca 'REVISAR PRECIO'."),
    ("Días de operación por mes", 26, "0", "DIAS_OPERACION", "Para venta diaria necesaria."),
    ("Efectivo inicial en banco (flujo)", 1500000, CRC, "EFECTIVO_INICIAL", "Saldo al inicio del flujo de caja."),
    ("Mes de inicio del flujo de caja", dt.date(2026, 11, 1), MONTH, "INICIO_FLUJO", "Día 1 del primer mes proyectado."),
    ("Crecimiento mensual de ventas (flujo)", 0.02, PCT, "CRECIMIENTO", "Supuesto de proyección."),
    ("Meses de gastos fijos como reserva mínima", 1, "0.0", "MESES_RESERVA", "Si el saldo baja de esto, el flujo alerta."),
]
for i, (lab, val, fmt, nm, note) in enumerate(params):
    r = 6 + i
    style(cf.cell(row=r, column=2, value=lab), F_BODY)
    v = cf.cell(row=r, column=3, value=val)
    style(v, F_IN, FILL_IN, fmt, Alignment(horizontal="right"))
    style(cf.cell(row=r, column=4, value=note), F_SUB, align=WRAP)
    name(nm, f"CONFIG!$C${r}")

LISTS = [
    ("F", "Categorías de insumo", ["Fruta", "Lácteos", "Abarrotes", "Toppings", "Jarabes y salsas", "Empaque", "Limpieza", "Otros"]),
    ("G", "Unidades base", ["g", "ml", "u"]),
    ("H", "Categorías de gasto", ["Alquiler", "Electricidad", "Agua", "Internet y teléfono", "Gas",
                                  "Publicidad y redes", "Servicios profesionales", "Mantenimiento y reparaciones",
                                  "Limpieza", "Patentes, permisos y seguros", "Comisiones tarjetas/plataformas",
                                  "Transporte y combustible", "Otros gastos"]),
    ("I", "Medios de pago", ["Efectivo", "SINPE Móvil", "Tarjeta", "Plataforma", "Transferencia"]),
    ("J", "Movimientos de inventario", ["Merma", "Vencimiento", "Daño", "Consumo interno", "Cortesía",
                                        "Ajuste (-) faltante", "Ajuste (+) sobrante"]),
    ("K", "Canales de venta", ["Local", "Para llevar", "WhatsApp", "Delivery"]),
    ("L", "Forma de pago compras", ["Contado", "Crédito"]),
    ("M", "Estado de pago", ["Pagado", "Pendiente"]),
    ("N", "Tipo de gasto", ["Fijo", "Variable"]),
    ("O", "Categorías de producto", ["Fresas con crema", "Especiales", "Extras", "Postres", "Bebidas"]),
]
LIST_ROWS = 15  # filas disponibles por lista (6..20)
LISTREF = {}
cf["F4"] = "LISTAS DESPLEGABLES (puede agregar opciones en las celdas azules vacías)"
cf["F4"].font = F_SEC
for col, head, items in LISTS:
    cf.column_dimensions[col].width = 24
    h = cf[f"{col}5"]
    h.value = head
    style(h, F_HDR, FILL_HDR_IN, align=CENTER)
    for k in range(LIST_ROWS):
        c = cf[f"{col}{6 + k}"]
        style(c, F_IN, FILL_IN)
        if k < len(items):
            c.value = items[k]
    LISTREF[head] = f"CONFIG!${col}$6:${col}${5 + LIST_ROWS}"
cf.row_dimensions[5].height = 30
cf.freeze_panes = "A6"

# ==========================================================================
# 3. INSUMOS (catálogo + inventario)
# ==========================================================================
ins = wb.create_sheet("INSUMOS")
title(ins, "INSUMOS E INVENTARIO",
      "Catálogo de materias primas y empaques. Todas las cantidades en UNIDAD BASE (g, ml o u). "
      "Stock actual = inicial + compras − consumo por ventas (recetas) ± mermas/ajustes.",
      "Costo vigente = último costo de COMPRAS (si existe) o el precio de catálogo. "
      "Costo real = costo ÷ % aprovechamiento (cáscara, hojas, residuos).")
A = "$A{r}"
ins_cols = [
    dict(h="Código", w=9, kind="in"),
    dict(h="Insumo", w=28, kind="in"),
    dict(h="Categoría", w=14, kind="in"),
    dict(h="Unidad base", w=8, kind="in", note="g, ml o u. Todo el libro usa esta unidad para este insumo."),
    dict(h="Presentación de compra", w=16, kind="in"),
    dict(h="Contenido por presentación (unid. base)", w=13, kind="in", fmt=NUM1,
         note="Ej.: caja de 1 kg = 1000 g; paquete de 50 vasos = 50 u."),
    dict(h="Precio presentación sin IVA (catálogo)", w=13, kind="in", fmt=CRC),
    dict(h="% IVA en compra", w=9, kind="in", fmt=PCT, note="Canasta básica 1%, general 13%."),
    dict(h="% aprovecha-miento", w=11, kind="in", fmt=PCT,
         note="Parte realmente utilizable. Fresas ~85%, mango ~65%. 100% si no hay desperdicio."),
    dict(h="Proveedor principal", w=16, kind="in"),
    dict(h="Stock mínimo (unid. base)", w=11, kind="in", fmt=NUM),
    dict(h="Stock inicial (unid. base)", w=11, kind="in", fmt=NUM,
         note="Inventario contado el día que empieza a usar el libro."),
    # M..
    dict(h="Costo catálogo por unid. base", w=11, kind="calc", fmt=CRC2,
         f=blank_guard(A, "IFERROR($G{r}/$F{r},0)")),
    dict(h="Último costo de compra por unid. base", w=11, kind="calc", fmt=CRC2,
         f=blank_guard(A, f"IFERROR(SUMIFS({R('COMPRAS','N')},{R('COMPRAS','D')},$A{{r}},{R('COMPRAS','R')},1)"
                          f"/COUNTIFS({R('COMPRAS','D')},$A{{r}},{R('COMPRAS','R')},1),0)")),
    dict(h="Costo vigente por unid. base", w=11, kind="calc", fmt=CRC2,
         f=blank_guard(A, "IF($N{r}>0,$N{r},$M{r})")),
    dict(h="Costo REAL por unid. base (con aprovech.)", w=12, kind="calc", fmt=CRC2,
         f=blank_guard(A, "IF(N($I{r})>0,$O{r}/$I{r},$O{r})")),
    dict(h="Entradas (compras)", w=11, kind="calc", fmt=NUM,
         f=blank_guard(A, f"SUMIFS({R('COMPRAS','I')},{R('COMPRAS','D')},$A{{r}})")),
    dict(h="Consumo por ventas (recetas)", w=11, kind="calc", fmt=NUM,
         f=blank_guard(A, f"SUMIFS({R('RECETAS','L')},{R('RECETAS','C')},$A{{r}})")),
    dict(h="Mermas y ajustes (±)", w=10, kind="calc", fmt=NUM,
         f=blank_guard(A, f"SUMIFS({R('MERMAS','J')},{R('MERMAS','B')},$A{{r}})")),
    dict(h="STOCK TEÓRICO ACTUAL", w=12, kind="calc", fmt=NUM,
         f=blank_guard(A, "N($L{r})+$Q{r}-$R{r}+$S{r}")),
    dict(h="Valor del inventario", w=13, kind="calc", fmt=CRC,
         f=blank_guard(A, "$T{r}*$O{r}")),
    dict(h="Conteo físico (unid. base)", w=11, kind="in", fmt=NUM,
         note="Escriba lo que contó en bodega. Déjelo vacío si no contó."),
    dict(h="Diferencia (físico − teórico)", w=11, kind="calc", fmt=NUM,
         f='IF(OR($A{r}="",$V{r}=""),"",$V{r}-$T{r})'),
    dict(h="Valor de la diferencia", w=12, kind="calc", fmt=CRC,
         f='IF($W{r}="","",$W{r}*$O{r})'),
    dict(h="Estado", w=12, kind="calc",
         f=blank_guard(A, 'IF($T{r}<=0,"AGOTADO",IF($T{r}<=N($K{r}),"REORDENAR","OK"))')),
    dict(h="Sugerido a comprar (presentaciones)", w=12, kind="calc", fmt=NUM,
         f=blank_guard(A, 'IF($T{r}<=N($K{r}),ROUNDUP((2*N($K{r})-$T{r})/$F{r},0),0)'),
         note="Lleva el stock al doble del mínimo."),
    dict(h="aux alerta", w=6, kind="calc",
         f='IF(OR($Y{r}="REORDENAR",$Y{r}="AGOTADO"),ROW(),"")'),
]
insumos = [
    ("I001", "Fresas frescas", "Fruta", "g", "Caja 5 kg", 5000, 14000, 0.01, 0.85, "Productor / feria", 5000, 12000),
    ("I002", "Mango maduro", "Fruta", "g", "Kilo", 1000, 1200, 0.01, 0.65, "Productor / feria", 2000, 5000),
    ("I003", "Pulpa de maracuyá", "Fruta", "g", "Bolsa 1 kg", 1000, 2800, 0.01, 1, "Distribuidora", 1000, 3000),
    ("I004", "Crema dulce", "Lácteos", "ml", "Caja 1 L", 1000, 2500, 0.01, 1, "Distribuidora lácteos", 4000, 10000),
    ("I005", "Leche condensada", "Lácteos", "g", "Lata 397 g", 397, 1250, 0.13, 0.97, "Mayorista", 3970, 7940),
    ("I006", "Esencia de vainilla", "Abarrotes", "ml", "Frasco 250 ml", 250, 1800, 0.13, 1, "Mayorista", 250, 500),
    ("I007", "Azúcar", "Abarrotes", "g", "Bolsa 2 kg", 2000, 1700, 0.01, 1, "Mayorista", 2000, 4000),
    ("I008", "Galleta Oreo triturada", "Toppings", "g", "Paquete 432 g", 432, 2600, 0.13, 0.95, "Mayorista", 864, 2160),
    ("I009", "Brownie (topping)", "Toppings", "g", "Bandeja 1 kg", 1000, 5500, 0.13, 0.95, "Panadería", 500, 1000),
    ("I010", "Marshmallows", "Toppings", "g", "Bolsa 500 g", 500, 2200, 0.13, 1, "Mayorista", 500, 1000),
    ("I011", "Gomitas", "Toppings", "g", "Bolsa 1 kg", 1000, 3800, 0.13, 1, "Mayorista", 500, 1000),
    ("I012", "Granola", "Toppings", "g", "Bolsa 1 kg", 1000, 4200, 0.13, 1, "Mayorista", 500, 1000),
    ("I013", "Maní picado", "Toppings", "g", "Bolsa 1 kg", 1000, 3600, 0.13, 1, "Mayorista", 500, 1000),
    ("I014", "Almendras fileteadas", "Toppings", "g", "Bolsa 1 kg", 1000, 9500, 0.13, 1, "Mayorista", 300, 1000),
    ("I015", "Arándanos deshidratados", "Toppings", "g", "Bolsa 1 kg", 1000, 8500, 0.13, 1, "Mayorista", 300, 1000),
    ("I016", "Queso rallado", "Lácteos", "g", "Bolsa 1 kg", 1000, 6500, 0.01, 1, "Distribuidora lácteos", 300, 1000),
    ("I017", "Jarabe de chocolate", "Jarabes y salsas", "g", "Botella 623 g", 623, 2900, 0.13, 0.95, "Mayorista", 623, 1246),
    ("I018", "Nutella", "Jarabes y salsas", "g", "Frasco 750 g", 750, 6800, 0.13, 0.95, "Mayorista", 750, 1500),
    ("I019", "Dulce de leche", "Jarabes y salsas", "g", "Frasco 1 kg", 1000, 3200, 0.13, 0.97, "Mayorista", 1000, 2000),
    ("I020", "Vaso 8 oz con tapa", "Empaque", "u", "Paquete 50 u", 50, 3750, 0.13, 1, "Empaques CR", 100, 300),
    ("I021", "Vaso 12 oz con tapa", "Empaque", "u", "Paquete 50 u", 50, 4500, 0.13, 1, "Empaques CR", 100, 400),
    ("I022", "Vaso 16 oz con tapa", "Empaque", "u", "Paquete 50 u", 50, 5500, 0.13, 1, "Empaques CR", 100, 400),
    ("I023", "Vaso 24 oz con tapa", "Empaque", "u", "Paquete 50 u", 50, 7000, 0.13, 1, "Empaques CR", 50, 150),
    ("I024", "Cuchara plástica", "Empaque", "u", "Paquete 100 u", 100, 1500, 0.13, 1, "Empaques CR", 200, 1000),
    ("I025", "Servilleta", "Empaque", "u", "Paquete 500 u", 500, 2500, 0.13, 1, "Empaques CR", 500, 2000),
    ("I026", "Bolsa para llevar", "Empaque", "u", "Paquete 100 u", 100, 2000, 0.13, 1, "Empaques CR", 100, 300),
    ("I027", "Crema de pistacho", "Jarabes y salsas", "g", "Frasco 200 g", 200, 4500, 0.13, 0.95, "Importadora", 200, 600),
    ("I028", "Harina", "Abarrotes", "g", "Bolsa 1 kg", 1000, 1100, 0.01, 1, "Mayorista", 1000, 2000),
    ("I029", "Mantequilla", "Lácteos", "g", "Barra 500 g", 500, 2000, 0.01, 1, "Distribuidora lácteos", 500, 1000),
    ("I030", "Cacao en polvo", "Abarrotes", "g", "Bolsa 500 g", 500, 3500, 0.13, 1, "Mayorista", 250, 500),
    ("I031", "Huevos", "Abarrotes", "u", "Cartón 30 u", 30, 3000, 0.01, 1, "Granja", 30, 60),
]
register(ins, "INSUMOS", ins_cols, insumos)
ins.column_dimensions["AA"].hidden = True
dv_list(ins, LISTREF["Categorías de insumo"], f"C{FIRST}:C{LAST['INSUMOS']}")
dv_list(ins, LISTREF["Unidades base"], f"D{FIRST}:D{LAST['INSUMOS']}")
status_colors(ins, f"Y{FIRST}:Y{LAST['INSUMOS']}", f"$Y{FIRST}", ["AGOTADO"], ["REORDENAR"])
ins.conditional_formatting.add(
    f"W{FIRST}:X{LAST['INSUMOS']}",
    FormulaRule(formula=[f'AND(ISNUMBER($W{FIRST}),$W{FIRST}<0)'],
                font=Font(name=FONT, bold=True, color="C00000")))
# totales arriba
ins["T4"] = "Valor total inventario:"
ins["T4"].font = F_BOLD
ins["T4"].alignment = Alignment(horizontal="right")
ins["U4"] = f"=SUM(U{FIRST}:U{LAST['INSUMOS']})"
style(ins["U4"], F_BOLD, FILL_KEY, CRC)
ins["W4"] = "Faltantes ₡:"
ins["W4"].font = F_BOLD
ins["X4"] = f"=SUM(X{FIRST}:X{LAST['INSUMOS']})"
style(ins["X4"], F_BOLD, FILL_KEY, CRC)

# ==========================================================================
# 4. COMPRAS
# ==========================================================================
com = wb.create_sheet("COMPRAS")
title(com, "REGISTRO DE COMPRAS A PROVEEDORES",
      "Una línea por insumo de cada factura. La cantidad se escribe en PRESENTACIONES (cajas, paquetes, latas); "
      "el libro la convierte a unidad base y la suma al inventario.")
D = "$D{r}"
com_cols = [
    dict(h="Fecha", w=11, kind="in", fmt=DATE),
    dict(h="N° factura", w=11, kind="in"),
    dict(h="Proveedor", w=18, kind="in"),
    dict(h="Código insumo", w=9, kind="in"),
    dict(h="Insumo", w=24, kind="calc", f=blank_guard(D, f'IFERROR({LOOK("INSUMOS","B",D)},"¿Código?")')),
    dict(h="Presentación", w=14, kind="calc", f=blank_guard(D, f'IFERROR({LOOK("INSUMOS","E",D)},"")')),
    dict(h="Cantidad (presenta-ciones)", w=10, kind="in", fmt=NUM1),
    dict(h="Contenido por presentación", w=11, kind="calc", fmt=NUM1,
         f=blank_guard(D, f'IFERROR({LOOK("INSUMOS","F",D)},0)')),
    dict(h="Cantidad en unid. base", w=11, kind="calc", fmt=NUM, f=blank_guard(D, "N($G{r})*$H{r}")),
    dict(h="Subtotal sin IVA", w=12, kind="in", fmt=CRC),
    dict(h="% IVA", w=8, kind="calc", fmt=PCT, f=blank_guard(D, f'IFERROR({LOOK("INSUMOS","H",D)},IVA_VENTAS)')),
    dict(h="IVA", w=10, kind="calc", fmt=CRC, f=blank_guard(D, "N($J{r})*$K{r}")),
    dict(h="Total factura línea", w=12, kind="calc", fmt=CRC, f=blank_guard(D, "N($J{r})+$L{r}")),
    dict(h="Costo por unid. base", w=10, kind="calc", fmt=CRC2, f=blank_guard(D, "IFERROR(N($J{r})/$I{r},0)")),
    dict(h="Forma de pago", w=10, kind="in"),
    dict(h="Estado de pago", w=10, kind="in"),
    dict(h="Mes", w=10, kind="calc", fmt=MONTH, f=blank_guard("$A{r}", "DATE(YEAR($A{r}),MONTH($A{r}),1)")),
    dict(h="Última compra del insumo (1=sí)", w=9, kind="calc",
         f=f'IF(OR($A{{r}}="",$D{{r}}=""),"",IF(COUNTIFS($D${FIRST}:$D${LAST["COMPRAS"]},$D{{r}},'
           f'$A${FIRST}:$A${LAST["COMPRAS"]},">"&$A{{r}})=0,1,0))'),
]
compras = [
    (dt.date(2026, 10, 1), "FA-1021", "Productor / feria", "I001", None, None, 3, None, None, 42000, None, None, None, None, "Contado", "Pagado"),
    (dt.date(2026, 10, 1), "FA-1021", "Productor / feria", "I002", None, None, 5, None, None, 6500, None, None, None, None, "Contado", "Pagado"),
    (dt.date(2026, 10, 2), "D-5530", "Distribuidora lácteos", "I004", None, None, 12, None, None, 30600, None, None, None, None, "Crédito", "Pendiente"),
    (dt.date(2026, 10, 2), "M-8812", "Mayorista", "I005", None, None, 24, None, None, 30000, None, None, None, None, "Contado", "Pagado"),
    (dt.date(2026, 10, 2), "M-8812", "Mayorista", "I008", None, None, 4, None, None, 10400, None, None, None, None, "Contado", "Pagado"),
    (dt.date(2026, 10, 3), "E-0099", "Empaques CR", "I021", None, None, 4, None, None, 18000, None, None, None, None, "Crédito", "Pendiente"),
]
register(com, "COMPRAS", com_cols, compras)
dv_list(com, R("INSUMOS", "A"), f"D{FIRST}:D{LAST['COMPRAS']}")
dv_list(com, LISTREF["Forma de pago compras"], f"O{FIRST}:O{LAST['COMPRAS']}")
dv_list(com, LISTREF["Estado de pago"], f"P{FIRST}:P{LAST['COMPRAS']}")
status_colors(com, f"P{FIRST}:P{LAST['COMPRAS']}", f"$P{FIRST}", ["Pendiente"], good=("Pagado",))
com["L4"] = "Pendiente de pago:"
com["L4"].font = F_BOLD
com["M4"] = f'=SUMIFS(M{FIRST}:M{LAST["COMPRAS"]},P{FIRST}:P{LAST["COMPRAS"]},"Pendiente")'
style(com["M4"], F_BOLD, FILL_KEY, CRC)

# ==========================================================================
# 5. PRODUCTOS (se crea antes que RECETAS en el código; el orden de hojas se ajusta al final)
# ==========================================================================
pro = wb.create_sheet("PRODUCTOS")
title(pro, "PRODUCTOS DEL MENÚ: COSTEO Y PRECIOS",
      "Precio de venta CON IVA (el que ve el cliente). Costo variable viene de RECETAS. "
      "Precio sugerido = costo ÷ (1 − margen objetivo)  →  'margen sobre ventas', no recargo sobre costo.",
      "Contribución = precio neto − costo variable: lo que cada venta aporta para pagar fijos (alquiler, planilla…) y generar utilidad.")
P = "$A{r}"
pro_cols = [
    dict(h="Código", w=8, kind="in"),
    dict(h="Producto", w=32, kind="in"),
    dict(h="Categoría", w=14, kind="in"),
    dict(h="Precio venta CON IVA", w=11, kind="in", fmt=CRC),
    dict(h="% IVA", w=7, kind="in", fmt=PCT),
    dict(h="Precio neto (sin IVA)", w=11, kind="calc", fmt=CRC, f=blank_guard(P, "IFERROR($D{r}/(1+$E{r}),0)")),
    dict(h="Porciones por receta / lote", w=10, kind="in", fmt=NUM,
         note="1 si la receta es por porción. Ej.: lote de brownie rinde 20 porciones."),
    dict(h="Costo variable por porción", w=11, kind="calc", fmt=CRC,
         f=blank_guard(P, f"SUMIFS({R('RECETAS','J')},{R('RECETAS','A')},$A{{r}})")),
    dict(h="Contribu-ción unitaria", w=11, kind="calc", fmt=CRC, f=blank_guard(P, "$F{r}-$H{r}")),
    dict(h="Margen de contribución %", w=10, kind="calc", fmt=PCT, f=blank_guard(P, "IFERROR($I{r}/$F{r},0)")),
    dict(h="Food cost %", w=9, kind="calc", fmt=PCT, f=blank_guard(P, "IFERROR($H{r}/$F{r},0)")),
    dict(h="Margen objetivo %", w=9, kind="in", fmt=PCT),
    dict(h="Precio sugerido neto", w=11, kind="calc", fmt=CRC,
         f=blank_guard(P, "IFERROR($H{r}/(1-IF(N($L{r})>0,$L{r},MARGEN_OBJ)),0)")),
    dict(h="Precio sugerido CON IVA", w=11, kind="calc", fmt=CRC, f=blank_guard(P, "$M{r}*(1+$E{r})")),
    dict(h="Recargo sobre costo %", w=10, kind="calc", fmt=PCT, f=blank_guard(P, "IFERROR($I{r}/$H{r},0)")),
    dict(h="Ventas estimadas por mes (unid.)", w=11, kind="in", fmt=NUM,
         note="Para PUNTO_EQUILIBRIO (objetivo de ventas)."),
    dict(h="Unidades vendidas (registradas)", w=11, kind="calc", fmt=NUM,
         f=blank_guard(P, f"SUMIFS({R('VENTAS','E')},{R('VENTAS','C')},$A{{r}})")),
    dict(h="Estado", w=14, kind="calc",
         f=blank_guard(P, 'IF($H{r}=0,"SIN RECETA",IF($K{r}>FOODCOST_MAX,"REVISAR PRECIO","OK"))')),
    dict(h="Preparación / notas de la ficha técnica", w=60, kind="in"),
]
productos = [
    ("P001", "Fresas con crema Pequeño (8 oz)", "Fresas con crema", 1950, 0.13, None, 1, None, None, None, None, 0.65, None, None, None, 300, None, None,
     "1) Lavar y desinfectar fresas, quitar hojas y picar en cuartos. 2) Batir crema dulce con leche condensada y vainilla. "
     "3) Capas: fresas – crema – fresas. 4) 1 topping incluido. 5) Glaseado leche condensada. Tapar y servir con cuchara."),
    ("P002", "Fresas con crema Mediano (12 oz)", "Fresas con crema", 2950, 0.13, None, 1, None, None, None, None, 0.65, None, None, None, 450, None, None,
     "Igual al pequeño en vaso de 12 oz. Incluye 2 toppings."),
    ("P003", "Fresas con crema Grande (16 oz)", "Fresas con crema", 4950, 0.13, None, 1, None, None, None, None, 0.65, None, None, None, 250, None, None,
     "Vaso de 16 oz. Incluye 3 toppings."),
    ("P004", "Fresas con crema Extra grande (24 oz)", "Fresas con crema", 6950, 0.13, None, 1, None, None, None, None, 0.65, None, None, None, 80, None, None,
     "Vaso de 24 oz. Incluye 4 toppings y bolsa."),
    ("P005", "Maracumango", "Especiales", 4200, 0.13, None, 1, None, None, None, None, 0.65, None, None, None, 150, None, None,
     "Mango en cubos + pulpa de maracuyá + crema. Vaso 16 oz."),
    ("P006", "Topping extra", "Extras", 200, 0.13, None, 1, None, None, None, None, 0.65, None, None, None, 200, None, None,
     "Porción 15 g (Oreo, granola, maní, gomitas…). Costeado con Oreo."),
    ("P007", "Topping extra premium", "Extras", 300, 0.13, None, 1, None, None, None, None, 0.65, None, None, None, 80, None, None,
     "Porción 15 g (almendras, arándanos o queso). Costeado con almendras."),
    ("P008", "Jarabe premium (Nutella / pistacho)", "Extras", 400, 0.13, None, 1, None, None, None, None, 0.65, None, None, None, 120, None, None,
     "Porción 25 g. Costeado con Nutella."),
    ("P009", "Brownie de la casa (lote de 20)", "Postres", 1500, 0.13, None, 20, None, None, None, None, 0.65, None, None, None, 60, None, None,
     "Ejemplo de LOTE: la receta es para 20 porciones; el costo se divide entre 20. Horno 180 °C, 25 min."),
]
register(pro, "PRODUCTOS", pro_cols, productos)
dv_list(pro, LISTREF["Categorías de producto"], f"C{FIRST}:C{LAST['PRODUCTOS']}")
status_colors(pro, f"R{FIRST}:R{LAST['PRODUCTOS']}", f"$R{FIRST}", ["REVISAR PRECIO"], ["SIN RECETA"])
for r in range(FIRST, LAST["PRODUCTOS"] + 1):
    pro[f"S{r}"].alignment = Alignment(wrap_text=False)

# ==========================================================================
# 6. RECETAS
# ==========================================================================
rec = wb.create_sheet("RECETAS")
title(rec, "RECETAS (FICHAS DE COSTO POR INGREDIENTE)",
      "Una línea por ingrediente. Cantidad NETA usada por porción (o por lote si en PRODUCTOS indicó más porciones). "
      "Incluya empaque (vaso, tapa, cuchara, servilleta). La nómina fija NO va aquí.")
RA = "$A{r}"
ok = 'OR($A{r}="",$C{r}="")'
rec_cols = [
    dict(h="Código producto", w=9, kind="in"),
    dict(h="Producto", w=30, kind="calc", f=blank_guard(RA, f'IFERROR({LOOK("PRODUCTOS","B",RA)},"¿Código?")')),
    dict(h="Código insumo", w=9, kind="in"),
    dict(h="Insumo", w=24, kind="calc", f=blank_guard("$C{r}", f'IFERROR({LOOK("INSUMOS","B","$C{r}")},"¿Código?")')),
    dict(h="Unidad", w=7, kind="calc", f=blank_guard("$C{r}", f'IFERROR({LOOK("INSUMOS","D","$C{r}")},"")')),
    dict(h="Cantidad neta por receta", w=10, kind="in", fmt=NUM1),
    dict(h="Costo real por unidad", w=10, kind="calc", fmt=CRC2,
         f=f'IF({ok},"",IFERROR({LOOK("INSUMOS","P","$C{r}")},0))'),
    dict(h="Costo de la línea (receta)", w=11, kind="calc", fmt=CRC2, f=f'IF({ok},"",N($F{{r}})*$G{{r}})'),
    dict(h="Porciones de la receta", w=9, kind="calc", fmt=NUM,
         f=f'IF({ok},"",IFERROR(MAX(1,{LOOK("PRODUCTOS","G",RA)}),1))'),
    dict(h="Costo por porción", w=10, kind="calc", fmt=CRC2, f=f'IF({ok},"",$H{{r}}/$I{{r}})'),
    dict(h="Unidades vendidas del producto", w=10, kind="calc", fmt=NUM,
         f=f'IF({ok},"",SUMIFS({R("VENTAS","E")},{R("VENTAS","C")},$A{{r}}))'),
    dict(h="Consumo bruto de inventario", w=11, kind="calc", fmt=NUM1,
         f=f'IF({ok},"",IFERROR(N($F{{r}})/{LOOK("INSUMOS","I","$C{r}")},N($F{{r}}))/$I{{r}}*$K{{r}})',
         note="Cantidad neta ÷ % aprovechamiento ÷ porciones × unidades vendidas."),
    dict(h="N° línea", w=6, kind="calc", f=f'IF($A{{r}}="","",COUNTIF($A${FIRST}:$A{{r}},$A{{r}}))'),
    dict(h="Clave", w=10, kind="calc", f='IF($A{r}="","",$A{r}&"|"&$M{r})'),
]
recipes = {
    "P001": [("I001", 90), ("I004", 45), ("I005", 20), ("I006", 1), ("I008", 15), ("I020", 1), ("I024", 1), ("I025", 1)],
    "P002": [("I001", 140), ("I004", 65), ("I005", 30), ("I006", 1.5), ("I008", 15), ("I012", 15), ("I021", 1), ("I024", 1), ("I025", 1)],
    "P003": [("I001", 220), ("I004", 100), ("I005", 45), ("I006", 2), ("I008", 20), ("I010", 15), ("I011", 15), ("I022", 1), ("I024", 1), ("I025", 2)],
    "P004": [("I001", 320), ("I004", 150), ("I005", 65), ("I006", 3), ("I008", 25), ("I010", 20), ("I011", 20), ("I009", 30),
             ("I023", 1), ("I024", 1), ("I025", 2), ("I026", 1)],
    "P005": [("I002", 180), ("I003", 60), ("I004", 60), ("I005", 30), ("I022", 1), ("I024", 1), ("I025", 1)],
    "P006": [("I008", 15)],
    "P007": [("I014", 15)],
    "P008": [("I018", 25)],
    "P009": [("I028", 400), ("I029", 250), ("I007", 400), ("I030", 120), ("I031", 6), ("I006", 10), ("I025", 20)],
}
rec_rows = []
for p, lines in recipes.items():
    for code, q in lines:
        rec_rows.append((p, None, code, None, None, q))
register(rec, "RECETAS", rec_cols, rec_rows)
dv_list(rec, R("PRODUCTOS", "A"), f"A{FIRST}:A{LAST['RECETAS']}")
dv_list(rec, R("INSUMOS", "A"), f"C{FIRST}:C{LAST['RECETAS']}")
rec.column_dimensions["M"].hidden = True
rec.column_dimensions["N"].hidden = True

# ==========================================================================
# 7. FICHA TÉCNICA
# ==========================================================================
ft = wb.create_sheet("FICHA_TECNICA")
title(ft, "FICHA TÉCNICA DEL PRODUCTO",
      "Elija el código del producto en la celda azul. Se arma sola con la receta, costos y margen.")
for col, w in zip("ABCDEFGH", (3, 6, 30, 10, 8, 14, 14, 3)):
    ft.column_dimensions[col].width = w
ft["B4"] = "Producto:"
ft["B4"].font = F_BOLD
ft["C4"] = "P001"
style(ft["C4"], Font(name=FONT, size=12, bold=True, color="0000FF"), FILL_IN)
dv_list(ft, R("PRODUCTOS", "A"), "C4")
ft["D4"] = f'=IFERROR({LOOK("PRODUCTOS","B","$C$4")},"")'
ft["D4"].font = Font(name=FONT, size=13, bold=True, color=C_DARK)
info = [
    ("Categoría", f'IFERROR({LOOK("PRODUCTOS","C","$C$4")},"")', None),
    ("Porciones por receta", f'IFERROR({LOOK("PRODUCTOS","G","$C$4")},1)', NUM),
    ("Precio venta CON IVA", f'IFERROR({LOOK("PRODUCTOS","D","$C$4")},0)', CRC),
    ("Precio neto (sin IVA)", f'IFERROR({LOOK("PRODUCTOS","F","$C$4")},0)', CRC),
]
for i, (lab, f, fmt) in enumerate(info):
    r = 5 + i
    ft.cell(row=r, column=3, value=lab).font = F_BODY
    c = ft.cell(row=r, column=4, value="=" + f)
    style(c, F_LINK, FILL_CALC, fmt)
ft.merge_cells("D5:E5")
hdr = ["#", "Ingrediente / insumo", "Cantidad", "Unidad", "Costo unitario", "Costo total"]
for j, h in enumerate(hdr):
    style(ft.cell(row=10, column=2 + j, value=h), F_HDR, FILL_HDR, align=CENTER)
NLINES = 25
for n in range(1, NLINES + 1):
    r = 10 + n
    key = f'$C$4&"|"&{n}'
    m = f"MATCH({key},{R('RECETAS','N')},0)"
    vals = [
        (n, None, None),
        (f'=IFERROR(INDEX({R("RECETAS","D")},{m}),"")', None, None),
        (f'=IFERROR(INDEX({R("RECETAS","F")},{m}),"")', NUM1, None),
        (f'=IFERROR(INDEX({R("RECETAS","E")},{m}),"")', None, None),
        (f'=IFERROR(INDEX({R("RECETAS","G")},{m}),"")', CRC2, None),
        (f'=IFERROR(INDEX({R("RECETAS","H")},{m}),"")', CRC2, None),
    ]
    for j, (v, fmt, _) in enumerate(vals):
        c = ft.cell(row=r, column=2 + j, value=v)
        style(c, F_BODY, FILL_CALC, fmt)
tr = 10 + NLINES + 1
lines_tot = [
    ("Costo total de la receta", f"=SUM(G11:G{10 + NLINES})", CRC),
    ("Costo variable por porción", f"=IFERROR(G{tr}/D6,0)", CRC),
    ("Contribución por porción", f"=D8-G{tr + 1}", CRC),
    ("Margen de contribución %", f"=IFERROR(G{tr + 2}/D8,0)", PCT),
    ("Food cost %", f"=IFERROR(G{tr + 1}/D8,0)", PCT),
    ("Precio sugerido CON IVA (margen objetivo)", f'=IFERROR({LOOK("PRODUCTOS","N","$C$4")},0)', CRC),
    ("Estado", f'=IFERROR({LOOK("PRODUCTOS","R","$C$4")},"")', None),
]
for i, (lab, f, fmt) in enumerate(lines_tot):
    r = tr + i
    ft.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    c = ft.cell(row=r, column=2, value=lab)
    style(c, F_BOLD, FILL_TOT, align=Alignment(horizontal="right"))
    c = ft.cell(row=r, column=7, value=f)
    style(c, F_BOLD, FILL_KEY, fmt)
pr = tr + len(lines_tot) + 1
ft.cell(row=pr, column=2, value="Preparación / notas").font = F_SEC
ft.merge_cells(start_row=pr + 1, start_column=2, end_row=pr + 6, end_column=7)
c = ft.cell(row=pr + 1, column=2, value=f'=IFERROR({LOOK("PRODUCTOS","S","$C$4")}&"","")')
style(c, F_BODY, FILL_CALC, align=WRAP)
status_colors(ft, f"G{tr + 6}", f"$G${tr + 6}", ["REVISAR PRECIO"], ["SIN RECETA"])

# ==========================================================================
# 8. VENTAS
# ==========================================================================
ven = wb.create_sheet("VENTAS")
title(ven, "REGISTRO DE VENTAS",
      "Puede registrar cada ticket o, al cierre del día, el total vendido por producto y medio de pago (reporte del POS). "
      "Cada venta descuenta del inventario los insumos de su receta.")
C = "$C{r}"
ven_cols = [
    dict(h="Fecha", w=11, kind="in", fmt=DATE),
    dict(h="N° ticket / factura", w=11, kind="in"),
    dict(h="Código producto", w=9, kind="in"),
    dict(h="Producto", w=30, kind="calc", f=blank_guard(C, f'IFERROR({LOOK("PRODUCTOS","B",C)},"¿Código?")')),
    dict(h="Cantidad", w=8, kind="in", fmt=NUM),
    dict(h="Precio unitario CON IVA", w=11, kind="calc", fmt=CRC, f=blank_guard(C, f'IFERROR({LOOK("PRODUCTOS","D",C)},0)')),
    dict(h="Descuento (₡)", w=10, kind="in", fmt=CRC),
    dict(h="Venta total CON IVA", w=12, kind="calc", fmt=CRC, f=blank_guard(C, "N($E{r})*$F{r}-N($G{r})")),
    dict(h="IVA", w=10, kind="calc", fmt=CRC,
         f=blank_guard(C, f'$H{{r}}-$H{{r}}/(1+IFERROR({LOOK("PRODUCTOS","E",C)},IVA_VENTAS))')),
    dict(h="Venta neta", w=12, kind="calc", fmt=CRC, f=blank_guard(C, "$H{r}-$I{r}")),
    dict(h="Costo variable (receta)", w=12, kind="calc", fmt=CRC,
         f=blank_guard(C, f'N($E{{r}})*IFERROR({LOOK("PRODUCTOS","H",C)},0)')),
    dict(h="Contribución", w=12, kind="calc", fmt=CRC, f=blank_guard(C, "$J{r}-$K{r}")),
    dict(h="Medio de pago", w=12, kind="in"),
    dict(h="Canal", w=10, kind="in"),
    dict(h="Mes", w=10, kind="calc", fmt=MONTH, f=blank_guard("$A{r}", "DATE(YEAR($A{r}),MONTH($A{r}),1)")),
]
ventas = []
day_mix = [
    (dt.date(2026, 10, 3), [("P001", 12, "Efectivo"), ("P002", 18, "SINPE Móvil"), ("P003", 9, "Tarjeta"),
                            ("P004", 3, "SINPE Móvil"), ("P005", 6, "Efectivo"), ("P006", 8, "Efectivo"), ("P008", 5, "SINPE Móvil")]),
    (dt.date(2026, 10, 4), [("P001", 15, "Efectivo"), ("P002", 22, "SINPE Móvil"), ("P003", 11, "Tarjeta"),
                            ("P004", 4, "Tarjeta"), ("P005", 7, "SINPE Móvil"), ("P007", 4, "Efectivo"), ("P009", 6, "Efectivo")]),
]
for d, items in day_mix:
    for code, q, mp in items:
        ventas.append((d, "Cierre " + d.strftime("%d/%m"), code, None, q, None, None, None, None, None, None, None, mp, "Local"))
register(ven, "VENTAS", ven_cols, ventas)
dv_list(ven, R("PRODUCTOS", "A"), f"C{FIRST}:C{LAST['VENTAS']}")
dv_list(ven, LISTREF["Medios de pago"], f"M{FIRST}:M{LAST['VENTAS']}")
dv_list(ven, LISTREF["Canales de venta"], f"N{FIRST}:N{LAST['VENTAS']}")

# ==========================================================================
# 9. MERMAS Y AJUSTES
# ==========================================================================
mer = wb.create_sheet("MERMAS")
title(mer, "MERMAS, CORTESÍAS Y AJUSTES DE INVENTARIO",
      "Cantidad siempre POSITIVA en unidad base. Todo resta del inventario excepto 'Ajuste (+) sobrante'. "
      "Use los ajustes para cuadrar con el conteo físico.")
B = "$B{r}"
mer_cols = [
    dict(h="Fecha", w=11, kind="in", fmt=DATE),
    dict(h="Código insumo", w=9, kind="in"),
    dict(h="Insumo", w=24, kind="calc", f=blank_guard(B, f'IFERROR({LOOK("INSUMOS","B",B)},"¿Código?")')),
    dict(h="Unidad", w=7, kind="calc", f=blank_guard(B, f'IFERROR({LOOK("INSUMOS","D",B)},"")')),
    dict(h="Tipo de movimiento", w=18, kind="in"),
    dict(h="Cantidad (unid. base)", w=10, kind="in", fmt=NUM1),
    dict(h="Motivo / detalle", w=30, kind="in"),
    dict(h="Responsable", w=14, kind="in"),
    dict(h="Costo por unid. base", w=10, kind="calc", fmt=CRC2, f=blank_guard(B, f'IFERROR({LOOK("INSUMOS","O",B)},0)')),
    dict(h="Cantidad con signo", w=10, kind="calc", fmt=NUM1,
         f=blank_guard(B, 'IF($E{r}="Ajuste (+) sobrante",1,-1)*N($F{r})')),
    dict(h="Valor (₡, negativo = pérdida)", w=12, kind="calc", fmt=CRC, f=blank_guard(B, "$J{r}*$I{r}")),
    dict(h="Mes", w=10, kind="calc", fmt=MONTH, f=blank_guard("$A{r}", "DATE(YEAR($A{r}),MONTH($A{r}),1)")),
]
mermas = [
    (dt.date(2026, 10, 3), "I001", None, None, "Merma", 600, "Fresas golpeadas / pasadas", "Encargado"),
    (dt.date(2026, 10, 4), "I004", None, None, "Vencimiento", 500, "Caja abierta vencida", "Encargado"),
]
register(mer, "MERMAS", mer_cols, mermas)
dv_list(mer, R("INSUMOS", "A"), f"B{FIRST}:B{LAST['MERMAS']}")
dv_list(mer, LISTREF["Movimientos de inventario"], f"E{FIRST}:E{LAST['MERMAS']}")

# ==========================================================================
# 10. PLANILLA
# ==========================================================================
pla = wb.create_sheet("PLANILLA")
title(pla, "PLANILLA MENSUAL Y CARGAS SOCIALES",
      "Una línea por colaborador por mes. Tasas en CONFIG (verifique las vigentes de CCSS e INS). "
      "Costo total empresa = bruto + CCSS patronal + INS + provisiones (aguinaldo, vacaciones, cesantía).")
PB = "$B{r}"
pla_cols = [
    dict(h="Mes (día 1)", w=10, kind="in", fmt=MONTH),
    dict(h="Colaborador", w=20, kind="in"),
    dict(h="Puesto", w=14, kind="in"),
    dict(h="Salario base mensual", w=12, kind="in", fmt=CRC),
    dict(h="Horas extra (cantidad)", w=9, kind="in", fmt=NUM1),
    dict(h="Valor hora extra", w=10, kind="calc", fmt=CRC, f=blank_guard(PB, "N($D{r})/HORAS_MES*RECARGO_HE")),
    dict(h="Monto horas extra", w=11, kind="calc", fmt=CRC, f=blank_guard(PB, "N($E{r})*$F{r}")),
    dict(h="Bonos / comisiones / otros", w=11, kind="in", fmt=CRC),
    dict(h="SALARIO BRUTO", w=12, kind="calc", fmt=CRC, f=blank_guard(PB, "N($D{r})+$G{r}+N($H{r})")),
    dict(h="CCSS obrera (rebajo)", w=11, kind="calc", fmt=CRC, f=blank_guard(PB, "$I{r}*CCSS_OBRERO")),
    dict(h="Otras deducciones (adelantos, renta, embargos)", w=13, kind="in", fmt=CRC),
    dict(h="NETO A PAGAR", w=12, kind="calc", fmt=CRC, f=blank_guard(PB, "$I{r}-$J{r}-N($K{r})")),
    dict(h="CCSS patronal", w=11, kind="calc", fmt=CRC, f=blank_guard(PB, "$I{r}*CCSS_PATRONAL")),
    dict(h="INS riesgos del trabajo", w=10, kind="calc", fmt=CRC, f=blank_guard(PB, "$I{r}*INS_RT")),
    dict(h="Provisión aguinaldo", w=10, kind="calc", fmt=CRC, f=blank_guard(PB, "$I{r}*PROV_AGUINALDO")),
    dict(h="Provisión vacaciones", w=10, kind="calc", fmt=CRC, f=blank_guard(PB, "$I{r}*PROV_VACACIONES")),
    dict(h="Provisión cesantía", w=10, kind="calc", fmt=CRC, f=blank_guard(PB, "$I{r}*PROV_CESANTIA")),
    dict(h="COSTO TOTAL EMPRESA", w=13, kind="calc", fmt=CRC, f=blank_guard(PB, "$I{r}+SUM($M{r}:$Q{r})")),
    dict(h="Mes (normalizado)", w=10, kind="calc", fmt=MONTH,
         f='IF(OR($A{r}="",$B{r}=""),"",DATE(YEAR($A{r}),MONTH($A{r}),1))'),
]
oct1 = dt.date(2026, 10, 1)
planilla = [
    (oct1, "Colaborador 1", "Encargado(a)", 520000, 8, None, None, 0),
    (oct1, "Colaborador 2", "Dependiente", 400000, 4, None, None, 0),
    (oct1, "Colaborador 3", "Dependiente medio tiempo", 200000, 0, None, None, 0),
]
register(pla, "PLANILLA", pla_cols, planilla)
pla["H4"] = "Mes del DASHBOARD →"
pla["H4"].font = F_BOLD
pla["I4"] = f"=SUMIFS(I{FIRST}:I{LAST['PLANILLA']},S{FIRST}:S{LAST['PLANILLA']},MES_REPORTE)"
pla["L4"] = f"=SUMIFS(L{FIRST}:L{LAST['PLANILLA']},S{FIRST}:S{LAST['PLANILLA']},MES_REPORTE)"
pla["R4"] = f"=SUMIFS(R{FIRST}:R{LAST['PLANILLA']},S{FIRST}:S{LAST['PLANILLA']},MES_REPORTE)"
for c in ("I4", "L4", "R4"):
    style(pla[c], F_BOLD, FILL_KEY, CRC)

# ==========================================================================
# 11. GASTOS
# ==========================================================================
gas = wb.create_sheet("GASTOS")
title(gas, "GASTOS OPERATIVOS",
      "Todo gasto que NO sea insumo de receta ni planilla: alquiler, servicios, publicidad, contador, mantenimiento, comisiones, etc.")
GB = "$B{r}"
gas_cols = [
    dict(h="Fecha", w=11, kind="in", fmt=DATE),
    dict(h="Categoría", w=26, kind="in"),
    dict(h="Descripción", w=30, kind="in"),
    dict(h="Proveedor", w=18, kind="in"),
    dict(h="Fijo / Variable", w=10, kind="in"),
    dict(h="Monto sin IVA", w=12, kind="in", fmt=CRC),
    dict(h="% IVA", w=8, kind="in", fmt=PCT),
    dict(h="IVA", w=10, kind="calc", fmt=CRC, f=blank_guard(GB, "N($F{r})*N($G{r})")),
    dict(h="Total pagado", w=12, kind="calc", fmt=CRC, f=blank_guard(GB, "N($F{r})+$H{r}")),
    dict(h="Medio de pago", w=12, kind="in"),
    dict(h="Mes", w=10, kind="calc", fmt=MONTH, f=blank_guard("$A{r}", "DATE(YEAR($A{r}),MONTH($A{r}),1)")),
]
gastos = [
    (dt.date(2026, 10, 1), "Alquiler", "Alquiler local octubre", "Arrendador", "Fijo", 450000, 0.13, None, None, "Transferencia"),
    (dt.date(2026, 10, 5), "Electricidad", "Recibo eléctrico", "CNFL / ICE", "Fijo", 85000, 0.13, None, None, "Transferencia"),
    (dt.date(2026, 10, 5), "Agua", "Recibo de agua", "AyA", "Fijo", 20000, 0, None, None, "Transferencia"),
    (dt.date(2026, 10, 6), "Internet y teléfono", "Internet + celular WhatsApp", "Operador", "Fijo", 25000, 0.13, None, None, "Transferencia"),
    (dt.date(2026, 10, 7), "Publicidad y redes", "Pauta Instagram / TikTok", "Meta", "Variable", 40000, 0.13, None, None, "Tarjeta"),
    (dt.date(2026, 10, 10), "Servicios profesionales", "Contador mensual", "Contador", "Fijo", 60000, 0.13, None, None, "Transferencia"),
]
register(gas, "GASTOS", gas_cols, gastos)
dv_list(gas, LISTREF["Categorías de gasto"], f"B{FIRST}:B{LAST['GASTOS']}")
dv_list(gas, LISTREF["Tipo de gasto"], f"E{FIRST}:E{LAST['GASTOS']}")
dv_list(gas, LISTREF["Medios de pago"], f"J{FIRST}:J{LAST['GASTOS']}")

# ==========================================================================
# 12. CAJA DIARIA
# ==========================================================================
caj = wb.create_sheet("CAJA_DIARIA")
title(caj, "CIERRE DE CAJA DIARIO",
      "Ventas por medio de pago vienen de VENTAS (misma fecha). Efectivo esperado = fondo + ventas en efectivo − gastos pagados de caja − depósitos/retiros.")
CA = "$A{r}"
vh, va, vm = R("VENTAS", "H"), R("VENTAS", "A"), R("VENTAS", "M")
caj_cols = [
    dict(h="Fecha", w=11, kind="in", fmt=DATE),
    dict(h="Responsable", w=14, kind="in"),
    dict(h="Fondo inicial", w=11, kind="in", fmt=CRC),
    dict(h="Ventas efectivo", w=12, kind="calc", fmt=CRC, f=blank_guard(CA, f'SUMIFS({vh},{va},$A{{r}},{vm},"Efectivo")')),
    dict(h="Ventas SINPE Móvil", w=12, kind="calc", fmt=CRC, f=blank_guard(CA, f'SUMIFS({vh},{va},$A{{r}},{vm},"SINPE Móvil")')),
    dict(h="Ventas tarjeta", w=12, kind="calc", fmt=CRC, f=blank_guard(CA, f'SUMIFS({vh},{va},$A{{r}},{vm},"Tarjeta")')),
    dict(h="Otros medios", w=11, kind="calc", fmt=CRC, f=blank_guard(CA, f"SUMIFS({vh},{va},$A{{r}})-$D{{r}}-$E{{r}}-$F{{r}}")),
    dict(h="TOTAL VENTAS DÍA", w=12, kind="calc", fmt=CRC, f=blank_guard(CA, "SUM($D{r}:$G{r})")),
    dict(h="Gastos pagados de caja", w=11, kind="in", fmt=CRC),
    dict(h="Depósitos / retiros", w=11, kind="in", fmt=CRC),
    dict(h="Efectivo esperado", w=12, kind="calc", fmt=CRC, f=blank_guard(CA, "N($C{r})+$D{r}-N($I{r})-N($J{r})")),
    dict(h="Efectivo contado", w=12, kind="in", fmt=CRC),
    dict(h="Diferencia", w=11, kind="calc", fmt=CRC, f='IF(OR($A{r}="",$L{r}=""),"",$L{r}-$K{r})'),
    dict(h="Estado", w=11, kind="calc",
         f=blank_guard(CA, 'IF($L{r}="","PENDIENTE",IF(ABS($M{r})<1,"OK",IF($M{r}<0,"FALTANTE","SOBRANTE")))')),
    dict(h="Observaciones", w=30, kind="in"),
]
caja = [
    (dt.date(2026, 10, 3), "Encargado", 30000, None, None, None, None, None, 5000, 50000, None, 24200, None, None, "Faltan ₡1,000: revisar vuelto"),
    (dt.date(2026, 10, 4), "Encargado", 30000, None, None, None, None, None, 0, 60000, None, None, None, None, ""),
]
register(caj, "CAJA_DIARIA", caj_cols, caja)
status_colors(caj, f"N{FIRST}:N{LAST['CAJA_DIARIA']}", f"$N{FIRST}", ["FALTANTE"], ["SOBRANTE", "PENDIENTE"])

# ==========================================================================
# 13. ESTADO DE RESULTADOS
# ==========================================================================
er = wb.create_sheet("ESTADO_RESULTADOS")
title(er, "ESTADO DE RESULTADOS MENSUAL (REAL)",
      "Automático a partir de VENTAS, MERMAS, PLANILLA, GASTOS y COMPRAS. Costo de ventas = costo teórico de receta de lo vendido.")
er.column_dimensions["A"].width = 3
er.column_dimensions["B"].width = 40
MCOLS = [get_column_letter(3 + k) for k in range(12)]
for L in MCOLS + ["O"]:
    er.column_dimensions[L].width = 12
er["B4"] = "=\"Negocio: \"&NEGOCIO&\"   ·   Año \"&ANIO"
er["B4"].font = F_BOLD
style(er["B5"], F_HDR, FILL_HDR, align=CENTER)
er["B5"] = "Concepto (₡)"
for k, L in enumerate(MCOLS):
    c = er[f"{L}5"]
    c.value = f"=DATE(ANIO,{k + 1},1)"
    style(c, F_HDR, FILL_HDR, MONTH, CENTER)
style(er["O5"], F_HDR, FILL_HDR, align=CENTER)
er["O5"] = "TOTAL AÑO"

ER = {}  # etiqueta -> fila
rows = []  # (key, label, formula_template(L) or None, kind)
M = "{L}$5"
ven_m = lambda col: f"SUMIFS({R('VENTAS', col)},{R('VENTAS', 'O')},{M})"  # noqa: E731
rows += [
    ("sec", "INGRESOS", None, "sec"),
    ("vb", "Ventas brutas (con IVA)", ven_m("H"), "n"),
    ("iva", "(−) IVA cobrado", ven_m("I"), "n"),
    ("vn", "VENTAS NETAS", ven_m("J"), "tot"),
    ("cv", "(−) Costo de ventas (recetas)", ven_m("K"), "n"),
    ("ub", "UTILIDAD BRUTA (contribución)", "{L}{vn}-{L}{cv}", "tot"),
    ("mb", "Margen bruto %", "IFERROR({L}{ub}/{L}{vn},0)", "pct"),
    ("sec2", "GASTOS DE OPERACIÓN", None, "sec"),
    ("mer", "Mermas y ajustes de inventario",
     f"-SUMIFS({R('MERMAS','K')},{R('MERMAS','L')},{M})", "n"),
    ("pla", "Planilla (costo total empresa)",
     f"SUMIFS({R('PLANILLA','R')},{R('PLANILLA','S')},{M})", "n"),
]
gcats = [f"CONFIG!$H${6 + k}" for k in range(LIST_ROWS)]
for k, ref in enumerate(gcats):
    rows.append((f"g{k}", "=" + ref,
                 f'IF({ref}="",0,SUMIFS({R("GASTOS","F")},{R("GASTOS","K")},{M},{R("GASTOS","B")},{ref}))', "n"))
rows += [
    ("tgo", "TOTAL GASTOS DE OPERACIÓN", "SUM({L}{mer}:{L}{glast})", "tot"),
    ("uo", "UTILIDAD OPERATIVA (antes de impuestos)", "{L}{ub}-{L}{tgo}", "key"),
    ("mo", "Margen operativo %", "IFERROR({L}{uo}/{L}{vn},0)", "pct"),
    ("isr", "(−) Impuesto sobre la renta estimado", "IF({L}{uo}>0,{L}{uo}*RENTA,0)", "n"),
    ("un", "UTILIDAD NETA", "{L}{uo}-{L}{isr}", "key"),
    ("sec3", "INDICADORES", None, "sec"),
    ("fc", "Food cost % (costo ventas ÷ ventas netas)", "IFERROR({L}{cv}/{L}{vn},0)", "pct"),
    ("lc", "Labor cost % (planilla ÷ ventas netas)", "IFERROR({L}{pla}/{L}{vn},0)", "pct"),
    ("pc", "Prime cost % (costo ventas + planilla)", "IFERROR(({L}{cv}+{L}{pla})/{L}{vn},0)", "pct"),
    ("sec4", "IVA DEL MES (referencia para la declaración)", None, "sec"),
    ("ivd", "IVA cobrado en ventas (débito)", "{L}{iva}", "n"),
    ("ivc", "IVA pagado en compras y gastos (crédito)",
     f"SUMIFS({R('COMPRAS','L')},{R('COMPRAS','Q')},{M})+SUMIFS({R('GASTOS','H')},{R('GASTOS','K')},{M})", "n"),
    ("ivn", "IVA neto por pagar (estimado)", "{L}{ivd}-{L}{ivc}", "tot"),
    ("sec5", "OTROS DATOS", None, "sec"),
    ("cmp", "Compras de insumos del mes (sin IVA)", f"SUMIFS({R('COMPRAS','J')},{R('COMPRAS','Q')},{M})", "n"),
]
r0 = 6
for i, (key, *_rest) in enumerate(rows):
    ER[key] = r0 + i
ER["glast"] = ER[f"g{LIST_ROWS - 1}"]
for i, (key, label, f, kind) in enumerate(rows):
    r = r0 + i
    lc = er.cell(row=r, column=2, value=label)
    if kind == "sec":
        style(lc, F_SEC, FILL_SEC)
        for L in MCOLS + ["O"]:
            er[f"{L}{r}"].fill = FILL_SEC
        continue
    font = F_BOLD if kind in ("tot", "key") else F_BODY
    fill = FILL_KEY if kind == "key" else (FILL_TOT if kind == "tot" else None)
    style(lc, font, fill)
    fmt = PCT if kind == "pct" else CRC
    for L in MCOLS:
        expr = f.replace("{L}", L)
        for kk, rr in ER.items():
            expr = expr.replace("{" + kk + "}", str(rr))
        style(er[f"{L}{r}"], font, fill or FILL_CALC, fmt)
        er[f"{L}{r}"] = "=" + expr
    # total año
    if kind == "pct":
        expr = f.replace("{L}", "O")
        for kk, rr in ER.items():
            expr = expr.replace("{" + kk + "}", str(rr))
        er[f"O{r}"] = "=" + expr
    else:
        er[f"O{r}"] = f"=SUM(C{r}:N{r})"
    style(er[f"O{r}"], F_BOLD, fill or FILL_TOT, fmt)
# ocultar filas de categorías vacías no se puede dinámicamente; quedan en 0
er.freeze_panes = "C6"

# ==========================================================================
# 14. PUNTO DE EQUILIBRIO (presupuesto mensual, como la plantilla de clase)
# ==========================================================================
pe = wb.create_sheet("PUNTO_EQUILIBRIO")
title(pe, "PRESUPUESTO MENSUAL Y PUNTO DE EQUILIBRIO",
      "Sección 1 toma costo, precio y ventas estimadas de PRODUCTOS. Escriba el precio de la competencia y los gastos mensuales en azul.")
widths = dict(A=3, B=34, C=12, D=2, E=11, F=11, G=10, H=12, I=11, J=10, K=9, L=13, M=13, N=13)
for k, v in widths.items():
    pe.column_dimensions[k].width = v
pe["B5"] = "1) Estudio de precios y márgenes unitarios"
pe["B5"].font = F_SEC
heads = {"B": "Artículos o servicios", "C": "Precio de la competencia (neto)", "E": "Costo variable unitario (C)",
         "F": "Precio unitario neto (P)", "G": "Price gap vs competencia", "H": "Margen de contribución unitario",
         "I": "Margen de contribución %", "J": "Ventas unidades mes (Q)", "K": "Mix de ventas",
         "L": "Ventas (P×Q)", "M": "Costos (C×Q)", "N": "Utilidad bruta"}
for col, h in heads.items():
    c = pe[f"{col}6"]
    c.value = h
    style(c, F_HDR, FILL_HDR_IN if col == "C" else FILL_HDR, align=CENTER)
pe.row_dimensions[6].height = 45
NP = 20
p1, pN = 8, 8 + NP - 1
for n in range(NP):
    r = p1 + n
    pr_ = FIRST + n
    pb = f"$B{r}"
    cells = {
        "B": (f'=IF(PRODUCTOS!$B${pr_}="","",PRODUCTOS!$B${pr_})', F_LINK, FILL_CALC, None),
        "C": (None, F_IN, FILL_IN, CRC),
        "E": (f'=IF({pb}="","",PRODUCTOS!$H${pr_})', F_LINK, FILL_CALC, CRC),
        "F": (f'=IF({pb}="","",PRODUCTOS!$F${pr_})', F_LINK, FILL_CALC, CRC),
        "G": (f'=IF(OR({pb}="",N($C{r})=0),"",F{r}/C{r}-1)', F_BODY, FILL_CALC, PCT),
        "H": (f'=IF({pb}="","",F{r}-E{r})', F_BODY, FILL_CALC, CRC),
        "I": (f'=IF({pb}="","",IFERROR(H{r}/F{r},0))', F_BODY, FILL_CALC, PCT),
        "J": (f'=IF({pb}="","",N(PRODUCTOS!$P${pr_}))', F_LINK, FILL_CALC, NUM),
        "K": (f'=IF({pb}="","",IFERROR(J{r}/$J$7,0))', F_BODY, FILL_CALC, PCT),
        "L": (f'=IF({pb}="","",F{r}*J{r})', F_BODY, FILL_CALC, CRC),
        "M": (f'=IF({pb}="","",E{r}*J{r})', F_BODY, FILL_CALC, CRC),
        "N": (f'=IF({pb}="","",L{r}-M{r})', F_BODY, FILL_CALC, CRC),
    }
    for col, (v, font, fill, fmt) in cells.items():
        c = pe[f"{col}{r}"]
        if v is not None:
            c.value = v
        style(c, font, fill, fmt)
# fila de totales (como en la plantilla, arriba de los artículos)
tot = {
    "B": ("TOTAL / PROMEDIO PONDERADO", None),
    "E": (f"=IFERROR(M7/J7,0)", CRC), "F": (f"=IFERROR(L7/J7,0)", CRC),
    "H": (f"=F7-E7", CRC), "I": (f"=IFERROR(N7/L7,0)", PCT),
    "J": (f"=SUM(J{p1}:J{pN})", NUM), "K": (f"=SUM(K{p1}:K{pN})", PCT),
    "L": (f"=SUM(L{p1}:L{pN})", CRC), "M": (f"=SUM(M{p1}:M{pN})", CRC), "N": (f"=SUM(N{p1}:N{pN})", CRC),
}
for col in "BCEFGHIJKLMN":
    c = pe[f"{col}7"]
    v = tot.get(col)
    if v:
        c.value = v[0]
    style(c, F_BOLD, FILL_TOT, v[1] if v else None)
pe["C4"] = "Las ventas en esta hoja son NETAS (sin IVA)."
pe["C4"].font = F_SUB

s2 = pN + 3
pe[f"B{s2}"] = "2) Estado de resultados presupuestado (colones por mes)"
pe[f"B{s2}"].font = F_SEC
b = s2 + 1
lines = [
    ("Unidades de venta", "=J7", NUM, "calc"),
    ("Ventas (P×Q)", "=L7", CRC, "calc"),
    ("Costo variable directo (C×Q)", "=M7", CRC, "calc"),
    ("Utilidad bruta", "=N7", CRC, "tot"),
    ("% Margen bruto de contribución", f"=IFERROR(C{b + 3}/C{b + 1},0)", PCT, "key"),
]
for i, (lab, f, fmt, kind) in enumerate(lines):
    r = b + i
    style(pe.cell(row=r, column=2, value=lab), F_BOLD if kind != "calc" else F_BODY,
          FILL_KEY if kind == "key" else None)
    style(pe.cell(row=r, column=3, value=f), F_BOLD if kind != "calc" else F_BODY,
          FILL_KEY if kind == "key" else (FILL_TOT if kind == "tot" else FILL_CALC), fmt)
pe[f"E{b + 4}"] = f'="Por cada colón que vendo me quedan "&TEXT(C{b + 4}*100,"0")&" céntimos para pagar gastos operativos y generar utilidad."'
pe[f"E{b + 4}"].font = F_SUB

g0 = b + 6
pe[f"B{g0}"] = "Gastos operativos mensuales:"
style(pe[f"B{g0}"], F_BOLD)
gastos_pres = [
    ("Planilla (costo total empresa)",
     f"=SUMIFS({R('PLANILLA','R')},{R('PLANILLA','S')},MAX({R('PLANILLA','S')}))", "link",
     "Último mes registrado en PLANILLA. Puede sobrescribir."),
    ("Alquiler del local", 450000, "in", None),
    ("Electricidad", 85000, "in", None),
    ("Agua", 20000, "in", None),
    ("Internet y teléfono", 25000, "in", None),
    ("Gas", 10000, "in", None),
    ("Publicidad y redes", 40000, "in", None),
    ("Servicios profesionales (contador)", 60000, "in", None),
    ("Mantenimiento y limpieza", 25000, "in", None),
    ("Patentes, permisos y seguros", 20000, "in", None),
    ("Comisiones tarjetas / plataformas", 30000, "in", None),
    ("Depreciación de equipo (no es salida de efectivo)", 41667, "in",
     "Ej.: equipo de ₡2,500,000 / 60 meses."),
    ("Otros", 20000, "in", None),
]
gs, ge = g0 + 1, g0 + len(gastos_pres)
for i, (lab, v, kind, note) in enumerate(gastos_pres):
    r = gs + i
    style(pe.cell(row=r, column=2, value="   " + lab), F_BODY)
    c = pe.cell(row=r, column=3, value=v)
    style(c, F_LINK if kind == "link" else F_IN, FILL_CALC if kind == "link" else FILL_IN, CRC)
    if note:
        c.comment = Comment(note, "Control")
pe[f"C{g0}"] = f"=SUM(C{gs}:C{ge})"
style(pe[f"C{g0}"], F_BOLD, FILL_TOT, CRC)
PE_PLANILLA, PE_DEPREC, PE_TOTGASTOS = f"C{gs}", f"C{ge - 1}", f"C{g0}"
PE_VENTAS, PE_CV, PE_MARGEN = f"C{b + 1}", f"C{b + 2}", f"C{b + 4}"

k0 = ge + 2
res = [
    ("Utilidad neta antes de impuestos", f"=C{b + 3}-{PE_TOTGASTOS}", CRC, "key"),
    ("% Margen neto", f"=IFERROR(C{k0}/{PE_VENTAS},0)", PCT, "calc"),
    ("PUNTO DE EQUILIBRIO en ventas (₡ netos/mes)", f"=IFERROR({PE_TOTGASTOS}/{PE_MARGEN},0)", CRC, "key"),
    ("Punto de equilibrio en unidades/mes (mix actual)", f"=IFERROR({PE_TOTGASTOS}/H7,0)", NUM, "key"),
    ("Venta diaria necesaria para equilibrio (₡ netos)", f"=IFERROR(C{k0 + 2}/DIAS_OPERACION,0)", CRC, "calc"),
    ("Venta diaria necesaria CON IVA", f"=C{k0 + 4}*(1+IVA_VENTAS)", CRC, "calc"),
    ("Margen de seguridad % (cuánto pueden caer las ventas)", f"=IFERROR(({PE_VENTAS}-C{k0 + 2})/{PE_VENTAS},0)", PCT, "calc"),
    ("Utilidad meta mensual deseada", 300000, CRC, "in"),
    ("Ventas netas necesarias para la utilidad meta", f"=IFERROR(({PE_TOTGASTOS}+C{k0 + 7})/{PE_MARGEN},0)", CRC, "key"),
]
for i, (lab, f, fmt, kind) in enumerate(res):
    r = k0 + i
    style(pe.cell(row=r, column=2, value=lab), F_BOLD if kind == "key" else F_BODY,
          FILL_KEY if kind == "key" else None)
    c = pe.cell(row=r, column=3, value=f)
    if kind == "in":
        style(c, F_IN, FILL_IN, fmt)
    else:
        style(c, F_BOLD if kind == "key" else F_BODY, FILL_KEY if kind == "key" else FILL_CALC, fmt)
pe[f"E{k0 + 2}"] = "Fórmula: gastos operativos ÷ % margen de contribución"
pe[f"E{k0 + 3}"] = "Fórmula: gastos operativos ÷ contribución unitaria promedio ponderada"
for rr in (k0 + 2, k0 + 3):
    pe[f"E{rr}"].font = F_SUB

# Calculadoras (diapositivas "Precio y aporte" y "Recargo y margen")
c0 = k0 + len(res) + 2
pe[f"B{c0}"] = "3) Calculadora: precio y aporte (un solo producto)"
pe[f"B{c0}"].font = F_SEC
calc1 = [
    ("Costo variable por unidad", 1000, "in", CRC),
    ("Gastos fijos mensuales", 400000, "in", CRC),
    ("Meta de utilidad operativa", 200000, "in", CRC),
    ("Ventas previstas (unidades/mes)", 400, "in", NUM),
    ("Precio neto = Variable + (Fijos + Meta) ÷ Unidades", f"=C{c0 + 1}+IFERROR((C{c0 + 2}+C{c0 + 3})/C{c0 + 4},0)", "key", CRC),
    ("Contribución por unidad", f"=C{c0 + 5}-C{c0 + 1}", "calc", CRC),
    ("Margen de contribución %", f"=IFERROR(C{c0 + 6}/C{c0 + 5},0)", "calc", PCT),
    ("Aporte total de las unidades previstas", f"=C{c0 + 6}*C{c0 + 4}", "calc", CRC),
    ("Utilidad (aporte − fijos)", f"=C{c0 + 8}-C{c0 + 2}", "key", CRC),
    ("¿Y si solo vendo estas unidades?", 200, "in", NUM),
    ("Utilidad / (pérdida) con esas unidades", f"=C{c0 + 6}*C{c0 + 10}-C{c0 + 2}", "key", CRC),
]
for i, (lab, v, kind, fmt) in enumerate(calc1):
    r = c0 + 1 + i
    style(pe.cell(row=r, column=2, value=lab), F_BODY)
    c = pe.cell(row=r, column=3, value=v)
    style(c, F_IN if kind == "in" else F_BOLD, FILL_IN if kind == "in" else (FILL_KEY if kind == "key" else FILL_CALC), fmt)

c1 = c0 + len(calc1) + 2
pe[f"B{c1}"] = "4) Calculadora: recargo sobre costo vs margen sobre ventas"
pe[f"B{c1}"].font = F_SEC
calc2 = [
    ("Costo variable", 1000, "in", CRC),
    ("Recargo sobre el costo %", 0.5, "in", PCT),
    ("Precio con recargo = costo × (1 + recargo)", f"=C{c1 + 1}*(1+C{c1 + 2})", "key", CRC),
    ("Margen real sobre ventas %", f"=IFERROR((C{c1 + 3}-C{c1 + 1})/C{c1 + 3},0)", "calc", PCT),
    ("Margen objetivo sobre ventas %", 0.5, "in", PCT),
    ("Precio por margen = costo ÷ (1 − margen)", f"=IFERROR(C{c1 + 1}/(1-C{c1 + 5}),0)", "key", CRC),
    ("Contribución con ese precio", f"=C{c1 + 6}-C{c1 + 1}", "calc", CRC),
]
for i, (lab, v, kind, fmt) in enumerate(calc2):
    r = c1 + 1 + i
    style(pe.cell(row=r, column=2, value=lab), F_BODY)
    c = pe.cell(row=r, column=3, value=v)
    style(c, F_IN if kind == "in" else F_BOLD, FILL_IN if kind == "in" else (FILL_KEY if kind == "key" else FILL_CALC), fmt)
pe[f"E{c1 + 4}"] = "El mismo porcentaje da precios distintos. La contribución aún debe cubrir los fijos."
pe[f"E{c1 + 4}"].font = F_SUB
pe.freeze_panes = "C7"

# ==========================================================================
# 15. FLUJO DE CAJA 12 MESES
# ==========================================================================
fl = wb.create_sheet("FLUJO_CAJA")
title(fl, "FLUJO DE CAJA PROYECTADO – 12 MESES",
      "Plan de caja: caja inicial y cobros · compras y pagos previstos · reserva y acciones ante faltantes. "
      "Base: PUNTO_EQUILIBRIO y PLANILLA. Puede sobrescribir cualquier mes en las filas azules.")
fl.column_dimensions["A"].width = 3
fl.column_dimensions["B"].width = 42
FC = [get_column_letter(3 + k) for k in range(12)]
for L in FC + ["O"]:
    fl.column_dimensions[L].width = 12
style(fl["B5"], F_HDR, FILL_HDR, align=CENTER)
fl["B5"] = "Concepto (₡)"
for k, L in enumerate(FC):
    c = fl[f"{L}5"]
    c.value = "=INICIO_FLUJO" if k == 0 else f"=EDATE({FC[k - 1]}5,1)"
    style(c, F_HDR, FILL_HDR, MONTH, CENTER)
style(fl["O5"], F_HDR, FILL_HDR, align=CENTER)
fl["O5"] = "TOTAL"

pla_last = f"MAX({R('PLANILLA','S')})"
def pla_sum(col):
    return f"SUMIFS({R('PLANILLA', col)},{R('PLANILLA','S')},{pla_last})"

FL = {}
frows = [
    ("si", "SALDO INICIAL DE CAJA", None, "key"),
    ("sec1", "ENTRADAS", None, "sec"),
    ("vtas", "Ventas cobradas (con IVA)", None, "in"),
    ("prest", "Préstamos / aportes de socios", None, "in"),
    ("te", "TOTAL ENTRADAS", "SUM({L}{vtas}:{L}{prest})", "tot"),
    ("sec2", "SALIDAS (pagos previstos)", None, "sec"),
    ("comp", "Compras de insumos (con IVA)", f"{{L}}{{vtas}}/(1+IVA_VENTAS)*IFERROR(PUNTO_EQUILIBRIO!${PE_CV[0]}${PE_CV[1:]}/PUNTO_EQUILIBRIO!${PE_VENTAS[0]}${PE_VENTAS[1:]},0)*(1+IVA_COMPRAS_PROM)", "calc"),
    ("sal", "Salarios brutos", pla_sum("I"), "calc"),
    ("cargas", "CCSS patronal + INS", f"{pla_sum('M')}+{pla_sum('N')}", "calc"),
    ("agui", "Aguinaldo (se paga en diciembre)", f"IF(MONTH({{L}}$5)=12,12*{pla_sum('O')},0)", "calc"),
    ("gfij", "Gastos operativos (sin planilla ni depreciación)",
     f"PUNTO_EQUILIBRIO!${PE_TOTGASTOS[0]}${PE_TOTGASTOS[1:]}-PUNTO_EQUILIBRIO!${PE_PLANILLA[0]}${PE_PLANILLA[1:]}"
     f"-PUNTO_EQUILIBRIO!${PE_DEPREC[0]}${PE_DEPREC[1:]}", "calc"),
    ("ivap", "IVA neto a Hacienda (del mes anterior)", None, "calc_iva"),
    ("cuota", "Cuotas de préstamos", None, "in"),
    ("renta", "Impuesto sobre la renta / pagos parciales", None, "in"),
    ("otros", "Inversiones / compras de equipo", None, "in"),
    ("impr", "Imprevistos", "SUM({L}{comp}:{L}{otros})*IMPREVISTOS", "calc"),
    ("ts", "TOTAL SALIDAS", "SUM({L}{comp}:{L}{impr})", "tot"),
    ("sec3", "RESULTADO", None, "sec"),
    ("mov", "Movimiento del mes (entradas − salidas)", "{L}{te}-{L}{ts}", "tot"),
    ("sf", "SALDO FINAL DE CAJA", "{L}{si}+{L}{mov}", "key"),
    ("res", "Reserva mínima requerida",
     f"MESES_RESERVA*(PUNTO_EQUILIBRIO!${PE_TOTGASTOS[0]}${PE_TOTGASTOS[1:]}-PUNTO_EQUILIBRIO!${PE_DEPREC[0]}${PE_DEPREC[1:]})", "calc"),
    ("alerta", "Alerta", 'IF({L}{sf}<0,"FALTANTE",IF({L}{sf}<{L}{res},"BAJO RESERVA","OK"))', "txt"),
]
fr0 = 6
for i, row in enumerate(frows):
    FL[row[0]] = fr0 + i
for i, (key, label, f, kind) in enumerate(frows):
    r = fr0 + i
    lc = fl.cell(row=r, column=2, value=label)
    if kind == "sec":
        style(lc, F_SEC, FILL_SEC)
        for L in FC + ["O"]:
            fl[f"{L}{r}"].fill = FILL_SEC
        continue
    font = F_BOLD if kind in ("tot", "key") else F_BODY
    style(lc, font, FILL_KEY if kind == "key" else (FILL_TOT if kind == "tot" else None))
    for k, L in enumerate(FC):
        c = fl[f"{L}{r}"]
        if key == "si":
            c.value = "=EFECTIVO_INICIAL" if k == 0 else f"={FC[k - 1]}{FL['sf']}"
            style(c, F_BOLD, FILL_KEY, CRC)
            continue
        if key == "vtas":
            c.value = (f"={PE_VENTAS.replace('C', 'PUNTO_EQUILIBRIO!$C$')}*(1+IVA_VENTAS)" if k == 0
                       else f"={FC[k - 1]}{r}*(1+CRECIMIENTO)")
            style(c, F_IN, FILL_IN, CRC)
            continue
        if kind == "in":
            c.value = 0
            style(c, F_IN, FILL_IN, CRC)
            continue
        if kind == "calc_iva":
            vt = f"{{V}}{FL['vtas']}"
            cp = f"{{V}}{FL['comp']}"
            iv = (f"MAX(0,{vt}-{vt}/(1+IVA_VENTAS)-({cp}-{cp}/(1+IVA_COMPRAS_PROM)))")
            src = L if k == 0 else FC[k - 1]
            c.value = "=" + iv.replace("{V}", src)
            style(c, F_BODY, FILL_CALC, CRC)
            continue
        expr = f.replace("{L}", L)
        for kk, rr in FL.items():
            expr = expr.replace("{" + kk + "}", str(rr))
        c.value = "=" + expr
        style(c, font, FILL_KEY if kind == "key" else (FILL_TOT if kind == "tot" else FILL_CALC),
              None if kind == "txt" else CRC, CENTER if kind == "txt" else None)
    o = fl[f"O{r}"]
    if key in ("si",):
        o.value = f"=C{r}"
    elif key in ("sf",):
        o.value = f"=N{r}"
    elif key in ("res", "alerta"):
        o.value = None
    else:
        o.value = f"=SUM(C{r}:N{r})"
    style(o, F_BOLD, FILL_TOT, CRC)
# corrección: el aguinaldo se provisiona; la fila 'si' del primer mes usa EFECTIVO_INICIAL
status_colors(fl, f"C{FL['alerta']}:N{FL['alerta']}", f"C${FL['alerta']}", ["FALTANTE"], ["BAJO RESERVA"])
fl.freeze_panes = "C6"
a0 = FL["alerta"] + 2
fl[f"B{a0}"] = "ACCIONES ANTE FALTANTES (si aparece BAJO RESERVA o FALTANTE)"
fl[f"B{a0}"].font = F_SEC
acciones = [
    "Negociar crédito o plazo con proveedores (pasar compras de Contado a Crédito).",
    "Reducir inventario: comprar más seguido y en menor cantidad los perecederos.",
    "Empujar productos de mayor contribución (ver DASHBOARD) y combos.",
    "Posponer inversiones y gastos no esenciales.",
    "Línea de crédito revolutiva o aporte de socios ANTES del mes en rojo, no después.",
    "Separar mensualmente la provisión de aguinaldo, vacaciones e IVA en una cuenta aparte.",
]
for i, t in enumerate(acciones):
    style(fl.cell(row=a0 + 1 + i, column=2, value="• " + t), F_BODY, border=False)

# ==========================================================================
# 16. DASHBOARD
# ==========================================================================
db = wb.create_sheet("DASHBOARD")
title(db, "DASHBOARD DEL MES", "Cambie el mes en CONFIG → 'Mes a analizar en el DASHBOARD'.")
for col, w in zip("ABCDEFGHIJ", (3, 30, 15, 3, 30, 15, 3, 30, 15, 12)):
    db.column_dimensions[col].width = w
db["B4"] = "=NEGOCIO"
db["B4"].font = Font(name=FONT, size=13, bold=True, color=C_DARK)
db["C4"] = "=MES_REPORTE"
style(db["C4"], F_BOLD, FILL_KEY, MONTH, CENTER)
colidx = "MATCH(MES_REPORTE,ESTADO_RESULTADOS!$C$5:$N$5,0)"


def erv(key):
    return f"=IFERROR(INDEX(ESTADO_RESULTADOS!$C${ER[key]}:$N${ER[key]},{colidx}),0)"


kpis = [
    ("B", 6, "Ventas netas", erv("vn"), CRC),
    ("B", 7, "Costo de ventas", erv("cv"), CRC),
    ("B", 8, "Utilidad bruta", erv("ub"), CRC),
    ("B", 9, "Gastos de operación (incl. planilla)", erv("tgo"), CRC),
    ("B", 10, "UTILIDAD OPERATIVA", erv("uo"), CRC),
    ("E", 6, "Food cost %", erv("fc"), PCT),
    ("E", 7, "Labor cost %", erv("lc"), PCT),
    ("E", 8, "Prime cost %", erv("pc"), PCT),
    ("E", 9, "Margen operativo %", erv("mo"), PCT),
    ("E", 10, "IVA neto por pagar (estimado)", erv("ivn"), CRC),
    ("H", 6, "Valor del inventario hoy", f"=SUM({R('INSUMOS','U')})", CRC),
    ("H", 7, "Insumos por reordenar / agotados",
     f'=COUNTIF({R("INSUMOS","Y")},"REORDENAR")+COUNTIF({R("INSUMOS","Y")},"AGOTADO")', NUM),
    ("H", 8, "Diferencias de conteo físico (₡)", f"=SUM({R('INSUMOS','X')})", CRC),
    ("H", 9, "Compras pendientes de pago", f'=SUMIFS({R("COMPRAS","M")},{R("COMPRAS","P")},"Pendiente")', CRC),
    ("H", 10, "Productos con precio a revisar", f'=COUNTIF({R("PRODUCTOS","R")},"REVISAR PRECIO")', NUM),
]
for col, r, lab, f, fmt in kpis:
    lc = db[f"{col}{r}"]
    lc.value = lab
    style(lc, F_BOLD if "UTILIDAD" in lab else F_BODY, FILL_SEC)
    vc = db.cell(row=r, column=lc.column + 1, value=f)
    style(vc, Font(name=FONT, size=11, bold=True, color=C_DARK), FILL_KEY if "UTILIDAD" in lab else FILL_CALC, fmt)
db["B5"] = "Resultados"
db["E5"] = "Indicadores"
db["H5"] = "Control"
for c in ("B5", "E5", "H5"):
    db[c].font = F_SEC
db["F4"] = "Meta food cost ≤"
db["F4"].font = F_SUB
db["G4"] = "=FOODCOST_MAX"
db["G4"].number_format = PCT
db.conditional_formatting.add("F6", FormulaRule(formula=["$F$6>FOODCOST_MAX"],
                                                 fill=PatternFill("solid", fgColor="F8CBAD")))

# Tabla de productos del mes
t0 = 13
db[f"B{t0 - 1}"] = "Productos en el mes"
db[f"B{t0 - 1}"].font = F_SEC
th = ["Producto", "Unidades", "", "Venta neta", "Contribución", "", "% de la contribución total", "Margen %"]
cols_t = ["B", "C", "D", "E", "F", "G", "H", "I"]
for col, h in zip(cols_t, th):
    if h:
        style(db[f"{col}{t0}"], F_HDR, FILL_HDR, align=CENTER)
        db[f"{col}{t0}"] = h
NPD = 20
for n in range(NPD):
    r = t0 + 1 + n
    pr_ = FIRST + n
    code = f"PRODUCTOS!$A${pr_}"
    vm = f"{R('VENTAS','C')},{code},{R('VENTAS','O')},MES_REPORTE"
    vals = {
        "B": (f'=IF({code}="","",PRODUCTOS!$B${pr_})', None),
        "C": (f'=IF({code}="","",SUMIFS({R("VENTAS","E")},{vm}))', NUM),
        "E": (f'=IF({code}="","",SUMIFS({R("VENTAS","J")},{vm}))', CRC),
        "F": (f'=IF({code}="","",SUMIFS({R("VENTAS","L")},{vm}))', CRC),
        "H": (f'=IF({code}="","",IFERROR(F{r}/SUM($F${t0 + 1}:$F${t0 + NPD}),0))', PCT),
        "I": (f'=IF({code}="","",IFERROR(F{r}/E{r},0))', PCT),
    }
    for col, (v, fmt) in vals.items():
        c = db[f"{col}{r}"]
        c.value = v
        style(c, F_BODY, FILL_CALC, fmt)
from openpyxl.formatting.rule import DataBarRule  # noqa: E402
db.conditional_formatting.add(f"H{t0 + 1}:H{t0 + NPD}",
                              DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="5B9BD5"))

# Alertas de inventario
a0 = t0 + NPD + 3
db[f"B{a0 - 1}"] = "Insumos por reordenar (desde INSUMOS)"
db[f"B{a0 - 1}"].font = F_SEC
for col, h in zip(["B", "C", "E", "F", "H"], ["Insumo", "Stock actual", "Proveedor", "Sugerido comprar (presentaciones)", "Estado"]):
    style(db[f"{col}{a0}"], F_HDR, FILL_HDR, align=CENTER)
    db[f"{col}{a0}"] = h
NAL = 15
for n in range(1, NAL + 1):
    r = a0 + n
    rowref = f"SMALL({R('INSUMOS','AA')},{n})"
    def irow(col):
        return f'=IFERROR(INDEX(INSUMOS!${col}:${col},{rowref}),"")'
    for col, src, fmt in (("B", "B", None), ("C", "T", NUM), ("E", "J", None), ("F", "Z", NUM), ("H", "Y", None)):
        c = db[f"{col}{r}"]
        c.value = irow(src)
        style(c, F_BODY, FILL_CALC, fmt)
status_colors(db, f"H{a0 + 1}:H{a0 + NAL}", f"$H{a0 + 1}", ["AGOTADO"], ["REORDENAR"])

# Gráfico anual
ch = BarChart()
ch.type = "col"
ch.title = "Ventas netas vs utilidad operativa (año)"
ch.y_axis.title = "₡"
ch.height, ch.width = 8, 22
data = Reference(er, min_col=2, max_col=14, min_row=ER["vn"], max_row=ER["vn"])
ch.add_data(data, from_rows=True, titles_from_data=True)
data2 = Reference(er, min_col=2, max_col=14, min_row=ER["uo"], max_row=ER["uo"])
ch.add_data(data2, from_rows=True, titles_from_data=True)
ch.set_categories(Reference(er, min_col=3, max_col=14, min_row=5, max_row=5))
db.add_chart(ch, f"B{a0 + NAL + 3}")

# ==========================================================================
# Orden de hojas, colores de pestaña y guardado
# ==========================================================================
order = ["INICIO", "CONFIG", "DASHBOARD", "INSUMOS", "COMPRAS", "RECETAS", "FICHA_TECNICA", "PRODUCTOS",
         "VENTAS", "MERMAS", "PLANILLA", "GASTOS", "CAJA_DIARIA", "ESTADO_RESULTADOS",
         "PUNTO_EQUILIBRIO", "FLUJO_CAJA"]
wb._sheets = [wb[s] for s in order]
tabs = {"INICIO": C_DARK, "CONFIG": "7F7F7F", "DASHBOARD": "C00000",
        "INSUMOS": "2E8B8B", "COMPRAS": "2E8B8B", "RECETAS": "70AD47", "FICHA_TECNICA": "70AD47",
        "PRODUCTOS": "70AD47", "VENTAS": "2F75B5", "MERMAS": "2E8B8B", "PLANILLA": "ED7D31",
        "GASTOS": "ED7D31", "CAJA_DIARIA": "2F75B5", "ESTADO_RESULTADOS": "7030A0",
        "PUNTO_EQUILIBRIO": "7030A0", "FLUJO_CAJA": "7030A0"}
for s, colr in tabs.items():
    wb[s].sheet_properties.tabColor = colr
wb.active = 0
wb.save(OUT)
print("OK ->", OUT)
